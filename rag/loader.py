"""文档解析：把 pdf/docx/txt/md 转成纯文本与元信息。"""
from __future__ import annotations

from pathlib import Path


class LoadError(Exception):
    pass


def load_text(path: str | Path) -> dict:
    """解析单个文件，返回 {text, meta}。meta 含来源/页数或段数等。"""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in (".txt", ".md", ".markdown"):
        return _load_plain(path)
    if ext == ".pdf":
        return _load_pdf(path)
    if ext == ".docx":
        return _load_docx(path)
    raise LoadError(f"不支持的格式：{ext}（支持 .txt/.md/.pdf/.docx）")


def _load_plain(path: Path) -> dict:
    raw = path.read_bytes()
    text = _decode(raw)
    return {
        "text": text,
        "meta": {"source": "text", "chars": len(text)},
    }


def _load_pdf(path: Path) -> dict:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise LoadError("缺少 pypdf，请先安装依赖") from e

    reader = PdfReader(str(path))
    parts = []
    for i, page in enumerate(reader.pages):
        t = page.extract_text() or ""
        if t.strip():
            parts.append(t)
    text = "\n\n".join(parts).strip()
    if not text:
        raise LoadError("未从 PDF 提取到文字（可能是扫描版/图片型 PDF）")
    return {"text": text, "meta": {"source": "pdf", "pages": len(reader.pages), "chars": len(text)}}


def _load_docx(path: Path) -> dict:
    try:
        import docx
    except ImportError as e:  # pragma: no cover
        raise LoadError("缺少 python-docx，请先安装依赖") from e

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    # 顺带抽取表格文本，避免信息丢失
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise LoadError("文档中没有可提取的文字")
    return {"text": text, "meta": {"source": "docx", "paragraphs": len(parts), "chars": len(text)}}


def _decode(raw: bytes) -> str:
    for enc in ("utf-8", "gb18030", "utf-16"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")
